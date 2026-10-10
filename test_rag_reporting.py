"""Modular RAG report regression tests and a reusable, synthetic visual-QA fixture."""
import csv
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import load_workbook

import cancellation
from advanced_rag import CONFIG, index_corpus
from rag_reporting import (start_details, read_details, write_details, export_workbook,
                           workbook_tables, finalize_details, numeric)
from test_cancellation import Tracker

ROOT = Path(__file__).resolve().parent


def make_fixture(folder):
    """Synthetic measurements, not actual experiment results."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    start_details(folder, CONFIG, 'fixture-index-generation')
    data = read_details(folder)
    rows = []
    for i, (mode, status) in enumerate([('rag', 'ok'), ('base_model_fallback', 'ok'),
                                       ('rag', 'error'), ('rag', 'cancelled')], 1):
        query = f'Synthetic question {i}: how does this energy system work?'
        r = {'accepted': mode == 'rag', 'response_mode': mode, 'fallback_used': mode != 'rag',
             'reason': 'dense_similarity' if mode == 'rag' else 'insufficient_evidence',
             'best_dense_similarity': .72 if mode == 'rag' else .12,
             'selected_child_count': 4, 'parent_count': 2 if mode == 'rag' else 0,
             'context_characters': 1500 if mode == 'rag' else 0,
             'sources': [{'filename': 'synthetic-energy.txt', 'section': 'Energy systems'}] if mode == 'rag' else [],
             'timings': dict(query_embedding_s=.031 * i, faiss_search_s=.003, bm25_search_s=.002,
                             fusion_mmr_s=.004, evidence_parent_s=.001)}
        data['queries'].append({'execution_position': i, 'question_id': f'General:{i}',
                                'question': query, 'topic_sheet': 'General', 'question_type': 'Direct',
                                'status': status, 'retrieval': r})
        for phase, factor in [('retrieval', 1), ('generation', 10)]:
            missing = i == 4 and phase == 'generation'
            data['groups'].append({'phase': phase,
                'group': 'Question embedding, search and ranking' if phase == 'retrieval' else
                         'Context augmentation and response generation',
                'execution_position': i, 'model': 'fixture-model', 'run_id': f'fixture-{i}-{phase}',
                'status': status if phase == 'generation' else 'measured',
                'duration_s': .25 * i * factor, 'wall_time_s': .26 * i * factor,
                'emissions_kg': None if missing else .00001 * i * factor,
                'energy_kwh': None if missing else .00002 * i * factor,
                'cpu_energy_kwh': None if missing else .00001 * i * factor,
                'gpu_energy_kwh': 0.0, 'ram_energy_kwh': None if missing else .00001 * i * factor})
        rows.append({'question_id': f'General:{i}', 'question': query, 'answer': 'Synthetic answer',
                     'model_name': 'fixture-model', 'status': status, 'error': '' if status == 'ok' else 'Fixture interruption',
                     'topic_sheet': 'General', 'question_number': i, 'question_type': 'Direct',
                     'execution_position': i, 'ordering_mode': 'original', 'shuffle_seed': '',
                     'retrieval_latency_s': .25 * i, 'generation_latency_s': 2.5 * i,
                     'total_latency_s': 2.75 * i, 'retrieval_emissions_kg': .00001 * i,
                     'generation_emissions_kg': None if i == 4 else .0001 * i,
                     'total_emissions_kg': None if i == 4 else .00011 * i,
                     'total_emissions_g': None if i == 4 else .11 * i,
                     'energy_kwh': None if i == 4 else .00022 * i,
                     'retrieval_energy_kwh': .00002 * i, 'generation_energy_kwh': .0002 * i,
                     'cpu_energy_kwh': .00011 * i, 'gpu_energy_kwh': 0.0, 'ram_energy_kwh': .00011 * i,
                     'chunks': '[]', 'retrieved_k': 2 if mode == 'rag' else 0})
    data['status'] = 'cancelled'
    write_details(folder, data)
    pd.DataFrame(rows).to_csv(folder / 'answers.csv', index=False)
    (folder / 'run_status.json').write_text(json.dumps({'pipeline': 'rag', 'status': 'cancelled',
        'planned_queries': 5, 'attempted_queries': 4, 'completed_queries': 3,
        'cancelled_queries': 1, 'unstarted_queries': 1}))
    return data, rows


class RagReportingTests(unittest.TestCase):
    def test_workbook_preserves_raw_units_missing_zero_and_question_identity(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            data, _ = make_fixture(temp)
            workbook = load_workbook(export_workbook(temp), data_only=False)
            self.assertEqual(workbook.sheetnames, ['Group measurements', 'Step timings', 'Evidence', 'Configuration'])
            group = workbook['Group measurements']
            self.assertEqual(group.max_row, 9)
            self.assertAlmostEqual(group['M2'].value, .00001)
            self.assertEqual(group['P2'].value, 0)
            self.assertIsNone(group['M9'].value)
            self.assertEqual(group['C2'].value, data['queries'][0]['question'])
            self.assertEqual(workbook['Step timings'].max_row, 5)
            self.assertEqual(workbook['Evidence']['D3'].value, 'base_model_fallback')
            workbook.close()

    def test_text_cannot_become_excel_formula_and_nan_stays_blank(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            data, _ = make_fixture(temp)
            data['queries'][0]['question'] = '=1+1'
            data['groups'][0]['emissions_kg'] = float('nan')
            write_details(temp, data)
            self.assertIsNone(read_details(temp)['groups'][0]['emissions_kg'])
            workbook = load_workbook(export_workbook(temp, read_details(temp)), data_only=False)
            self.assertEqual(workbook['Group measurements']['C2'].data_type, 's')
            self.assertEqual(workbook['Group measurements']['C2'].value, '=1+1')
            self.assertIsNone(workbook['Group measurements']['M2'].value)
            workbook.close()

    def test_pdf_sections_and_history_are_not_invented(self):
        from generate_report import generate
        from pypdf import PdfReader
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            make_fixture(temp)
            text = '\n'.join(page.extract_text() for page in PdfReader(generate(Path(temp))).pages)
            for title in ('RAG evidence and fallback overview', 'Retrieval-step timings',
                          'Evidence selection by question', 'Evidence-backed versus fallback costs',
                          'Recorded retrieval configuration', 'fixture-index-generation'):
                self.assertIn(title, text)
            self.assertIn('insufficient evidence', text)
            self.assertIn('0.11 / 1', text)  # Only successful RAG question, not the error/cancelled costs.
            (Path(temp) / 'rag_details.json').unlink()
            text = '\n'.join(page.extract_text() for page in PdfReader(generate(Path(temp))).pages)
            self.assertIn('Detailed RAG decisions and timings were not recorded', text)
            self.assertNotIn('Retrieval-step timings', text)

    def test_indexing_two_groups_sum_and_saved_validation_are_measured(self):
        import numpy as np
        class Embedder:
            def encode(self, texts, **kwargs):
                return np.array([[1, .2] for _ in texts], dtype='float32')
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            base = Path(temp)
            corpus, output = base / 'corpus', base / 'report'
            corpus.mkdir()
            (corpus / 'source.txt').write_text('# Title\nAn energy measurement example.')
            trackers = []
            def factory(*args, **kwargs):
                t = Tracker(output, 'emissions.csv', kwargs['project_name'])
                trackers.append(t)
                return t
            @cancellation.cancellable_run
            def build():
                index_corpus(corpus, base / 'index', output, {'.txt'}, lambda p, c: p.read_text(),
                             Mock, lambda name: Embedder(), factory)
            build()
            self.assertEqual(len(trackers), 2)
            self.assertTrue(all(t.stops == 1 for t in trackers))
            details = read_details(output)
            self.assertEqual([g['phase'] for g in details['groups']],
                             ['document_preparation', 'embedding_indexing_storage'])
            self.assertTrue(details['index_generation'])
            self.assertTrue((output / 'modular_emissions.xlsx').exists())
            summary = pd.read_csv(output / 'summary.csv').iloc[0]
            self.assertAlmostEqual(summary['index_emissions_kg'], .002)
            self.assertAlmostEqual(summary['total_energy_kwh'], .006)
            self.assertEqual(workbook_tables(details).keys(), {'Group measurements', 'Configuration'})

    def test_export_failure_does_not_discard_durable_measurements(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            make_fixture(temp)
            with patch('rag_reporting.export_workbook', side_effect=PermissionError('Workbook open in Excel')):
                self.assertIsNone(finalize_details(temp, 'cancelled'))
            self.assertEqual(read_details(temp)['status'], 'cancelled')
            self.assertTrue(export_workbook(temp).is_file())

    def test_details_are_matched_by_position_and_exact_question(self):
        from rag_pdf import matching_details
        rows = [{'question': 'A', 'execution_position': 1}]
        details = {'queries': [{'question': 'B', 'execution_position': 1, 'retrieval': {'response_mode': 'rag'}}]}
        self.assertEqual(matching_details(rows, details), [None])
        self.assertIsNone(numeric('nan'))
        self.assertEqual(numeric(0), 0)


if __name__ == '__main__':
    unittest.main()
