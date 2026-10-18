"""Export five runs as six question-aligned Excel tables (no inference).

Run: python export_run_tables.py
Optional: python export_run_tables.py PATH1 PATH2 PATH3 PATH4 PATH5 --output result.xlsx
Output is a static analysis snapshot; rerun after source changes.
"""
import argparse
import csv
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from question_order import BANK, load_bank

ROOT = Path(__file__).resolve().parent
IDENTIFIERS = ['Original question serial', 'Question type - topic sheet']
METRICS = ['Execution position', 'CPU energy (kWh)', 'GPU energy (kWh)',
           'RAM energy (kWh)', 'Calculated CPU emissions (g CO2eq)',
           'Calculated GPU emissions (g CO2eq)', 'Calculated RAM emissions (g CO2eq)',
           'Latency (s)', 'Input tokens', 'Output tokens',
           'Reconstructed input tokens', 'Reconstructed saved-answer tokens']


def number(value, context):
    """Keep unavailable values blank, preserve zero, reject invalid measurements."""
    if value is None or str(value).strip().lower() in ('', 'nan', 'none', 'null'):
        return None
    result = float(value)
    if not math.isfinite(result) or result < 0:
        raise ValueError(f'{context}: invalid measurement {value!r}')
    return result


def read_run(folder, questions):
    path = folder / 'answers.csv'
    with path.open(encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle)
        required = {'question', 'execution_position', 'status', 'latency_s',
                    'cpu_energy_kwh', 'gpu_energy_kwh', 'ram_energy_kwh',
                    'energy_kwh', 'emissions_kg'}
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f'{path}: missing columns {sorted(missing)}')
        rows = list(reader)
    mapped = {}
    for row in rows:
        text = row['question']  # Exact identity: no trimming or case conversion.
        if text in mapped:
            raise ValueError(f'{folder.name}: duplicate question text {text[:80]!r}')
        mapped[text] = row
    expected = set(questions)
    if set(mapped) != expected:
        raise ValueError(f'{folder.name}: question set differs from workbook '
                         f'({len(expected-set(mapped))} missing, '
                         f'{len(set(mapped)-expected)} extra)')
    positions, result = [], []
    for text in questions:
        row = mapped[text]
        if row['status'] != 'ok':
            raise ValueError(f'{folder.name}: question has status {row["status"]!r}; '
                             'this export requires successful runs')
        position = number(row['execution_position'], folder.name)
        if position is None or not position.is_integer():
            raise ValueError(f'{folder.name}: invalid execution position')
        positions.append(int(position))
        energies = [number(row[f'{part}_energy_kwh'], f'{folder.name}/{part}')
                    for part in ('cpu', 'gpu', 'ram')]
        total = number(row['energy_kwh'], folder.name)
        emissions = number(row['emissions_kg'], folder.name)
        if total is not None and all(value is not None for value in energies):
            if not math.isclose(sum(energies), total, rel_tol=1e-8, abs_tol=1e-14):
                raise ValueError(f'{folder.name}: component energies do not sum to total')
        if total == 0 and emissions not in (None, 0):
            raise ValueError(f'{folder.name}: positive emissions with zero energy')
        factor = emissions * 1000 / total if total and emissions is not None else None
        derived = [value * factor if value is not None and factor is not None
                   else (0.0 if value == 0 and total == 0 and emissions == 0 else None)
                   for value in energies]
        result.append([int(position), *energies, *derived,
                       number(row['latency_s'], folder.name),
                       number(row.get('input_tokens'), folder.name),
                       number(row.get('output_tokens'), folder.name),
                       number(row.get('reconstructed_input_tokens'), folder.name),
                       number(row.get('reconstructed_output_tokens'), folder.name)])
    if sorted(positions) != list(range(1, len(questions) + 1)):
        raise ValueError(f'{folder.name}: execution positions must be a permutation of 1..N')
    return result


def add_table(workbook, name, headers, rows, notes, table_number):
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    end = len(rows) + 1
    # Plain cells avoid Excel/preview engines applying a default coloured table style.
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{end}'
    sheet.freeze_panes = 'C2'
    sheet.sheet_view.showGridLines = True
    sheet.row_dimensions[1].height = 62
    for column, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(column)].width = 39 if column == 2 else 23
        for row in range(1, end + 1):
            cell = sheet.cell(row, column)
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Workbook labels must never execute as formulas.
            if row == 1:
                cell.alignment = Alignment(wrap_text=True)
            elif column > 2:
                cell.number_format = ('0' if 'Execution position' in header or 'tokens' in header.lower() else
                                      '0.000' if 'Latency' in header else '0.000000000E+00')
        if column == 1:
            for row in range(2, end + 1):
                sheet.cell(row, column).number_format = '0'
    for index, note in enumerate(notes, end + 3):
        sheet.cell(index, 1, note)
    sheet.print_title_rows = '1:1'
    return sheet


def export(folders, output, bank_path=BANK):
    folders = [Path(folder).resolve() for folder in folders]
    output = Path(output).resolve()
    if len(folders) != 5 or len(set(folders)) != 5:
        raise ValueError('Choose five different run folders')
    names = [folder.name for folder in folders]
    if len(set(name.lower() for name in names)) != 5 or any(
        len(name) > 31 or any(c in name for c in '[]:*?/\\') or
        name.lower() == 'comprehensive' or name.startswith("'") or name.endswith("'")
        for name in names
    ):
        raise ValueError('Run folder names must be distinct valid Excel sheet names')
    if output.suffix.lower() != '.xlsx':
        raise ValueError('Output must have an .xlsx extension')
    if output.exists():
        raise FileExistsError(f'Output already exists; choose another --output path: {output}')
    bank = load_bank(bank_path)
    questions = bank['Question'].tolist()
    if len(set(questions)) != len(questions):
        raise ValueError('Questionbank contains duplicate question text; matching is ambiguous')
    labels = [[i, f'{row.question_type} - {row.topic_sheet}']
              for i, row in enumerate(bank.itertuples(index=False), 1)]
    runs = [read_run(folder, questions) for folder in folders]
    wb = Workbook()
    wb.remove(wb.active)
    notes = [
        'Rows follow the original questionbank sheet and row order. Questions match by exact text.',
        'Calculated component emissions = component kWh x (recorded total kg CO2eq / recorded total kWh) x 1000.',
        'Component emissions are allocated estimates, not separately measured emissions. Blank means unavailable; zero is retained.',
        'Static snapshot calculated by export_run_tables.py. Rerun to refresh after source changes.',
        'Reconstructed tokens recount saved text with current local tokenizer/template assets; they are not historical Ollama counters.',
        f'Question source: {Path(bank_path).resolve()}',
    ]
    headers = IDENTIFIERS + [f'{name}: {metric}' for metric in METRICS for name in names]
    combined = [label + [run[i][m] for m in range(len(METRICS)) for run in runs]
                for i, label in enumerate(labels)]
    add_table(wb, 'Comprehensive', headers, combined,
              notes + [f'Run source: {folder / "answers.csv"}' for folder in folders], 1)
    for index, (folder, rows) in enumerate(zip(folders, runs), 2):
        add_table(wb, folder.name, IDENTIFIERS + METRICS,
                  [label + row for label, row in zip(labels, rows)],
                  notes + [f'Run source: {folder / "answers.csv"}'], index)
    output.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation protects existing exports, including a concurrent writer.
    with output.open('xb') as handle:
        wb.save(handle)
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('folders', nargs='*', type=Path)
    parser.add_argument('--output', type=Path, default=ROOT / 'obsv' / 'run_tables.xlsx')
    parser.add_argument('--questionbank', type=Path, default=BANK)
    args = parser.parse_args()
    folders = args.folders or [ROOT / 'emissions_reports' / f'e{i}' for i in range(1, 6)]
    try:
        path = export(folders, args.output, args.questionbank)
    except (ValueError, OSError) as exc:
        parser.exit(1, f'Cannot export: {exc}\n')
    print(f'Created six tables: {path}')


if __name__ == '__main__':
    main()
