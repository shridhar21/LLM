import csv
import json
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(r"C:\Users\Asus\Desktop\llmtest")
SOURCE = ROOT / "emissions_reports"
OUTPUT = ROOT / "obsv"
RUNS = ["e1", "e2", "e3", "e4", "e5"]

MEASUREMENTS = {
    "cpu_energy_kwh",
    "ram_energy_kwh",
    "gpu_energy_kwh",
    "emissions_kg",
    "latency_s",
    "measurement_duration_s",
}

# These are redundant measurement columns in the original answers.csv.
DROP = {"emissions_g", "energy_kwh"}

INTEGER_COLUMNS = {
    "question_number",
    "execution_position",
    "workbook_position",
    "shuffle_seed",
    "conversation_turn",
}


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return reader.fieldnames or [], list(reader)


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def number(value):
    if value is None or str(value).strip().lower() in {"", "nan", "none"}:
        return None

    result = float(value)
    if not math.isfinite(result) or result < 0:
        return None

    return result


def excel_value(value, column):
    if column in MEASUREMENTS:
        return number(value)

    if column in INTEGER_COLUMNS:
        parsed = number(value)
        if parsed is None:
            return None
        if not parsed.is_integer():
            raise ValueError(f"{column} must be an integer: {value}")
        return int(parsed)

    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)

    return value


def write_sheet(workbook, name, columns, rows):
    sheet = workbook.create_sheet(name)
    sheet.append(columns)

    for row in rows:
        row_number = sheet.max_row + 1

        for position, column in enumerate(columns, 1):
            value = excel_value(row.get(column), column)

            if isinstance(value, str) and len(value) > 32767:
                raise ValueError(
                    f"{name}, row {row_number}: {column} exceeds "
                    "Excel's text limit. Text was not silently truncated."
                )

            cell = sheet.cell(row_number, position)
            cell.value = value

            # Treat source text as text, including strings starting with "=".
            if isinstance(value, str):
                cell.data_type = "s"

            cell.alignment = Alignment(vertical="top", wrap_text=True)

            if column in MEASUREMENTS:
                cell.number_format = "0.000000000E+00"
            elif column in INTEGER_COLUMNS:
                cell.number_format = "0"

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 45

    for position, column in enumerate(columns, 1):
        width = 65 if column in {"question", "answer", "value"} else 25
        sheet.column_dimensions[get_column_letter(position)].width = width


def prepare_run(run_name):
    folder = SOURCE / run_name

    columns, answers = read_csv(folder / "answers.csv")
    _, raw = read_csv(folder / "emissions.csv")
    metadata = read_json(folder / "measurement_metadata.json")
    ordering = read_json(folder / "question_order.json")

    required = {
        "question_id",
        "cpu_energy_kwh",
        "ram_energy_kwh",
        "gpu_energy_kwh",
        "emissions_kg",
        "latency_s",
    }

    if not required.issubset(columns):
        raise ValueError(f"{run_name}: missing {required - set(columns)}")

    if not answers:
        raise ValueError(f"{run_name}: answers.csv is empty")

    question_ids = [row["question_id"] for row in answers]
    if len(set(question_ids)) != len(question_ids):
        raise ValueError(f"{run_name}: duplicate question identities")

    raw_by_id = {}
    for row in raw:
        run_id = row["run_id"]
        if run_id in raw_by_id:
            raise ValueError(f"{run_name}: duplicate raw run ID {run_id}")
        raw_by_id[run_id] = row

    planned = {
        str(row["Question ID"]): row
        for row in ordering.get("questions", [])
    }

    output_rows = []

    for answer in answers:
        qid = answer["question_id"]
        prefix = f"LLM_only_q{qid}_"

        # Match by question identity and tracker run ID, never row position.
        matches = [
            session
            for session in metadata["sessions"]
            if session.get("project_name", "").startswith(prefix)
        ]

        if len(matches) != 1:
            raise ValueError(
                f"{run_name}/{qid}: expected one tracker, "
                f"found {len(matches)}"
            )

        session = matches[0]
        run_id = session["run_id"]

        if run_id not in raw_by_id:
            raise ValueError(f"{run_name}/{qid}: missing raw tracker row")

        row = {
            key: value
            for key, value in answer.items()
            if key not in DROP
        }

        row.update(
            measurement_duration_s=raw_by_id[run_id].get("duration"),
            shuffle_description=ordering.get("description"),
            conversation_order_policy=ordering.get(
                "conversation_order_policy"
            ),
            conversation_chain=planned.get(qid, {}).get(
                "conversation_chain"
            ),
            conversation_turn=planned.get(qid, {}).get(
                "conversation_turn"
            ),
            run_id=run_id,
            emissions_available=session.get("emissions_available"),
            tracker_stop_error=session.get("stop_error"),
            tracking_scope=session.get("effective_scope"),
            config_profile=session.get("config_profile"),
            hardware_evidence=session.get("hardware_evidence"),
        )

        output_rows.append(row)

    workbook = Workbook()
    workbook.remove(workbook.active)

    kept_columns = list(
        dict.fromkeys(column for row in output_rows for column in row)
    )
    write_sheet(workbook, "Questions", kept_columns, output_rows)

    # Keep complete configuration, hardware and measurement-quality evidence.
    sessions = metadata.get("sessions", [])
    session_columns = list(
        dict.fromkeys(column for session in sessions for column in session)
    )
    if session_columns:
        write_sheet(
            workbook,
            "Measurement metadata",
            session_columns,
            sessions,
        )

    details = {
        "source_folder": str(folder),
        "missing_measurements": (
            "Blank means unavailable; zero means recorded zero."
        ),
        "cpu_energy": (
            "Recorded CPU energy; inspect hardware evidence for backend."
        ),
        "ram_energy": (
            "CodeCarbon estimate, not directly metered RAM power."
        ),
        **{
            key: value
            for key, value in ordering.items()
            if key != "questions"
        },
    }

    if metadata.get("scope_note"):
        details["scope_note"] = metadata["scope_note"]

    write_sheet(
        workbook,
        "Run details",
        ["field", "value"],
        [
            {"field": key, "value": value}
            for key, value in details.items()
        ],
    )

    return workbook


def main():
    destinations = {
        name: OUTPUT / f"{name}.xlsx"
        for name in RUNS
    }

    existing = [
        path for path in destinations.values()
        if path.exists()
    ]
    if existing:
        raise FileExistsError(
            "Existing exports were not overwritten: "
            + ", ".join(str(path) for path in existing)
        )

    # Validate and prepare every run before writing any workbook.
    workbooks = {
        name: prepare_run(name)
        for name in RUNS
    }

    OUTPUT.mkdir(parents=True, exist_ok=True)

    for name, workbook in workbooks.items():
        destination = destinations[name]
        workbook.save(destination)
        print(f"Created {destination}")


if __name__ == "__main__":
    main()