"""Local dashboard contract checks; synthetic machine copies, no inference."""
import csv
import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import dashboard as d


class DashboardTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent.parent)
        self.root=Path(self.temp.name)
        self.old_output=d.OUTPUT;d.OUTPUT=self.root/'exports'
        self.client=d.app.test_client()
        self.bank=d.load_bank().iloc[:20]
        self.sources=[]
        for m in range(3):
            machine=self.root/f'pc{m}'
            self.sources.append({'name':f'PC {m+1}','path':str(machine)})
            for mode in ['original','questions']:
                folder=machine/f'arbitrary_{mode}';folder.mkdir(parents=True)
                rows=[]
                from question_order import arrange, save_order
                frame=arrange(self.bank,mode,42)
                for pos,(_,q) in enumerate(frame.iterrows(),1):
                    rows.append(dict(question_id=q['Question ID'],question=q['Question'],answer='fixture',
                                     model_name='fixture:3b',status='ok',error='',execution_position=pos,
                                     ordering_mode=mode,latency_s=2+m,emissions_kg=.001*(m+1),
                                     energy_kwh=.002,cpu_energy_kwh=.001,ram_energy_kwh=.001,gpu_energy_kwh=0))
                with (folder/'answers.csv').open('w',newline='',encoding='utf-8') as f:
                    w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
                with (folder/'summary.csv').open('w',newline='') as f:
                    w=csv.DictWriter(f,fieldnames=['num_queries']);w.writeheader();w.writerow({'num_queries':20})
                (folder/'run_status.json').write_text(json.dumps({'status':'completed'}))
                save_order(frame,folder)
        self.snap={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in self.root.rglob('*') if p.is_file()}
        response=self.client.post('/api/load',json={'machines':self.sources})
        self.assertEqual(response.status_code,200,response.data)
        loaded=response.get_json();self.dataset=loaded['dataset']
        self.ids=[r['id'] for machine in loaded['machines'] for r in machine['runs']]
        self.payload={'dataset':self.dataset,'ids':self.ids}

    def tearDown(self):
        for path,digest in self.snap.items():
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),digest)
        d.OUTPUT=self.old_output;d.CATALOGUES.clear();d.MACHINE_SETS.clear();d.ARTIFACTS.clear()
        self.temp.cleanup()

    def test_tables_matching_and_averages(self):
        response=self.client.post('/api/analyse',json=self.payload)
        self.assertEqual(response.status_code,200,response.data)
        result=response.get_json()
        self.assertEqual(len(result['questions']),120)
        self.assertEqual(len(result['followups']),5)
        self.assertEqual(len(result['averages']),60)
        self.assertTrue(all(r['mean_emissions_g']==int(r['machine'][-1]) for r in result['averages']))
        for row in result['followups']:
            gaps=[v for k,v in row.items() if k.endswith(' gap')]
            self.assertTrue(all(v is None for v in gaps) if row['turn']==1 else all(v==0 for v in gaps))

    def test_custom_graphs_and_all_exports(self):
        from pypdf import PdfReader
        for kind in ['execution','range','average','report_individual','report_comparison']:
            response=self.client.post('/api/graph',json={**self.payload,'kind':kind,'page':1})
            self.assertEqual(response.status_code,200,response.data)
            panels=response.get_json()['panels'];self.assertEqual(len(panels),3)
            for panel in panels:
                with self.client.get(panel['url']) as image_response:
                    self.assertEqual(image_response.status_code,200)
        for kind in ['csv','xlsx','comparison','custom','individual','side-by-side']:
            response=self.client.post('/api/export',json={**self.payload,'kind':kind,'run':self.ids[0],
                                                        'graph_kind':'range','page':1})
            self.assertEqual(response.status_code,200,response.data)
            result=response.get_json();token=result['url'].split('/')[-1];path=d.ARTIFACTS[token]
            self.assertGreater(path.stat().st_size,0)
            if path.suffix=='.pdf':self.assertGreater(len(PdfReader(path).pages),0)

    def test_reject_mixed_modes_changed_text_and_unknown_ids(self):
        item=d.CATALOGUES[self.dataset][self.ids[-1]]
        original=item['run']['metadata']['Pipeline']
        item['run']['metadata']['Pipeline']='RAG'
        self.assertEqual(self.client.post('/api/analyse',json=self.payload).status_code,400)
        item['run']['metadata']['Pipeline']=original
        q=next(iter(item['run']['questions']));record=item['run']['questions'].pop(q)
        item['run']['questions'][q+' changed']=record
        self.assertEqual(self.client.post('/api/analyse',json=self.payload).status_code,400)
        self.assertEqual(self.client.post('/api/analyse',json={**self.payload,'ids':['unknown']}).status_code,400)

    def test_loopback_and_origin_restrictions(self):
        self.assertEqual(self.client.get('/',headers={'Host':'example.com'}).status_code,403)
        self.assertEqual(self.client.post('/api/load',json={},headers={'Origin':'https://example.com'}).status_code,403)
        self.assertEqual(self.client.get('/artifact/not-real').status_code,404)

    def test_rag_phases_and_missing_machine_panel(self):
        # Convert only these synthetic fixtures; no real experiment data is changed.
        for path in self.root.glob('pc*/*/answers.csv'):
            with path.open(newline='',encoding='utf-8') as f:rows=list(csv.DictReader(f))
            for row in rows:
                row['total_emissions_kg']=row.pop('emissions_kg')
                row['total_latency_s']=row.pop('latency_s')
                row.update(retrieval_emissions_kg=.0002,generation_emissions_kg=float(row['total_emissions_kg'])-.0002,
                           retrieval_latency_s=.1,generation_latency_s=float(row['total_latency_s'])-.1,
                           retrieval_energy_kwh=.0005,generation_energy_kwh=.0015)
            with path.open('w',newline='',encoding='utf-8') as f:
                writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
        self.snap={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in self.snap}
        response=self.client.post('/api/load',json={'machines':self.sources})
        loaded=response.get_json()
        payload={'dataset':loaded['dataset'],'ids':[r['id'] for m in loaded['machines'] for r in m['runs']]}
        response=self.client.post('/api/graph',json={**payload,'kind':'phases','page':1})
        self.assertEqual(response.status_code,200,response.data)
        self.assertTrue(all('url' in p for p in response.get_json()['panels']))
        response=self.client.post('/api/graph',json={**payload,'ids':payload['ids'][:1],'kind':'phases','page':1})
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(sum('missing' in p for p in response.get_json()['panels']),2)
        explore=self.explorer_payload(averaging='schemes')
        explore.update(dataset=loaded['dataset'],pipeline='RAG')
        for selection,machine in zip(explore['selections'],loaded['machines']):
            selection['runs']=[r['id'] for r in machine['runs']]
        response=self.client.post('/api/explore',json=explore)
        self.assertEqual(response.status_code,200,response.data)
        for row in response.get_json()['rows']:
            self.assertAlmostEqual(row['retrieval_emissions_g'],.2)
            self.assertAlmostEqual(row['retrieval_emissions_g']+row['generation_emissions_g'],row['emissions_g'])
            self.assertAlmostEqual(row['retrieval_energy_wh'],.5)

    def test_source_changes_require_reload(self):
        path=self.root/'pc0'/'arbitrary_original'/'answers.csv'
        before=path.read_bytes()
        try:
            path.write_bytes(before+b'\n')
            response=self.client.post('/api/analyse',json=self.payload)
            self.assertEqual(response.status_code,400)
            self.assertIn('Reload',response.get_json()['error'])
        finally:path.write_bytes(before)

    def test_optional_machine_folders_and_validation(self):
        for count in (1,2,3):
            with self.subTest(machines=count):
                response=self.client.post('/api/load',json={'machines':self.sources[:count]+[{'path':'  '}]})
                self.assertEqual(response.status_code,200,response.data)
                loaded=response.get_json()
                self.assertEqual(len(loaded['machines']),count)
                ids=[r['id'] for m in loaded['machines'] for r in m['runs']]
                body={'dataset':loaded['dataset'],'ids':ids}
                self.assertEqual(self.client.post('/api/analyse',json=body).status_code,200)
                # Panel slots correspond only to loaded machines, even with no selected match.
                fake=[{'machine':m['index'],'label':m['name'],'missing':'test'} for m in loaded['machines']]
                with patch.object(d,'custom_files') as render:
                    render.side_effect=ValueError('render not needed')
                    with patch.object(d,'selected',return_value=([],self.bank)):
                        panels=d.make_graphs({**body,'kind':'execution'})
                    self.assertEqual(len(panels),count)
                    self.assertTrue(all('missing' in p for p in panels))
                from pypdf import PdfReader
                for layout in ('vertical','horizontal'):
                    with patch.object(d,'make_graphs',return_value=fake):
                        response=self.client.post('/api/export',json={**body,'kind':'side-by-side','layout':layout})
                    self.assertEqual(response.status_code,200,response.data)
                    path=d.ARTIFACTS[response.get_json()['url'].split('/')[-1]]
                    pages=PdfReader(path).pages
                    self.assertEqual(len(pages),count if layout=='vertical' else 1)
                    self.assertEqual(float(pages[0].mediabox.width),1584*(count if layout=='horizontal' else 1))
        for roots in ([],[{'path':' '}],self.sources+[self.sources[0]],self.sources[:1]*2):
            self.assertEqual(self.client.post('/api/load',json={'machines':roots}).status_code,400)

    def explorer_payload(self, count=3, averaging='none'):
        return dict(dataset=self.dataset,pipeline='Normal LLM',averaging=averaging,compare=True,
                    selections=[dict(machine=m,questions={'mode':'serials','serials':'1, 13-17'},
                                     models=['fixture:3b'],schemes=['original','questions'],
                                     runs=[id_ for id_ in self.ids if d.CATALOGUES[self.dataset][id_]['entry']['machine']==m])
                                for m in range(count)])

    def test_explorer_subset_and_averages_stay_within_model_and_machine(self):
        body=self.explorer_payload(averaging='schemes')
        response=self.client.post('/api/explore',json=body)
        self.assertEqual(response.status_code,200,response.data)
        result=response.get_json()
        self.assertEqual(len(result['raw']),36)
        self.assertEqual(len(result['rows']),18)
        self.assertTrue(all(r['emissions_g']==r['machine_index']+1 for r in result['rows']))
        self.assertTrue(all(r['scheme_count']==2 and r['run_count']==2 and r['position'] is None for r in result['rows']))
        # A second model must never be pooled into the first model's mean.
        rows=[dict(result['raw'][0]),dict(result['raw'][1])]
        rows[1].update(question=rows[0]['question'],original_question=rows[0]['original_question'],
                       model='another-model',emissions_g=100)
        means=d.analysis.average_rows(rows)
        self.assertEqual(sorted(r['emissions_g'] for r in means),[1,100])

    def test_explorer_two_models_from_real_fixture_files(self):
        import shutil
        for machine in range(3):
            for mode in ('original','questions'):
                source=self.root/f'pc{machine}'/f'arbitrary_{mode}'
                target=self.root/f'pc{machine}'/f'model_b_{mode}'
                shutil.copytree(source,target)
                path=target/'answers.csv'
                with path.open(newline='',encoding='utf-8') as f:rows=list(csv.DictReader(f))
                for row in rows:row.update(model_name='model-b',emissions_kg=str(.01*(machine+1)))
                with path.open('w',newline='',encoding='utf-8') as f:
                    writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
        response=self.client.post('/api/load',json={'machines':self.sources}).get_json()
        body=self.explorer_payload(averaging='schemes');body['dataset']=response['dataset']
        for selection,machine in zip(body['selections'],response['machines']):
            selection['models']=['fixture:3b','model-b'];selection['runs']=[r['id'] for r in machine['runs']]
        response=self.client.post('/api/explore',json=body)
        self.assertEqual(response.status_code,200,response.data)
        result=response.get_json()
        self.assertEqual(len(result['rows']),36)
        for row in result['rows']:
            self.assertEqual(row['emissions_g'],(row['machine_index']+1)*(10 if row['model']=='model-b' else 1))

    def test_equal_scheme_weights_missing_values_and_zero(self):
        data=self.client.post('/api/explore',json=self.explorer_payload(1)).get_json()
        base=data['raw'][0]
        rows=[dict(base,scheme='original',emissions_g=0),dict(base,scheme='original',emissions_g=4),
              dict(base,scheme='questions',emissions_g=10)]
        mean=d.analysis.average_rows(rows)[0]
        self.assertEqual(mean['emissions_g'],6)  # mean(mean(0,4),10), not mean(0,4,10)
        self.assertEqual(mean['run_count'],3)
        rows[0]['emissions_g']=None
        mean=d.analysis.average_rows(rows)[0]
        self.assertIsNone(mean['emissions_g'])
        self.assertEqual(mean['emissions_g_valid_runs'],2)

    def test_token_counts_average_per_question_model_and_machine(self):
        body=self.explorer_payload(1)
        raw=self.client.post('/api/explore',json=body).get_json()['raw']
        self.assertTrue(all(r['input_tokens'] is None and r['output_tokens'] is None for r in raw))
        rows=[dict(raw[0],input_tokens=100,output_tokens=0,scheme='original'),
              dict(raw[0],input_tokens=200,output_tokens=40,scheme='questions')]
        average=d.analysis.average_rows(rows)[0]
        self.assertEqual(average['input_tokens'],150)
        self.assertEqual(average['output_tokens'],20)

    def test_explorer_selection_methods_and_true_followup_predecessor(self):
        questions=d.analysis.question_records(d.load_bank())
        available=set(range(1,21))
        selected=d.analysis.choose_questions({'mode':'sheets','sheets':[questions[0]['topic']],'subset':'followups'},available,questions)
        self.assertEqual([q['serial'] for q in selected],[13,14,15,16,17])
        by_type=d.analysis.choose_questions({'mode':'types','types':['__followups__']},available,questions)
        self.assertEqual(selected,by_type)
        body=self.explorer_payload(1);body['selections'][0]['questions']={'mode':'serials','serials':'15'}
        response=self.client.post('/api/explore',json=body)
        self.assertEqual(response.status_code,200,response.data)
        self.assertTrue(all(r['gap']==0 for r in response.get_json()['followups']))
        for text in ('0','4-2','9999','1,no'):
            with self.assertRaises(ValueError):d.analysis.choose_questions({'mode':'serials','serials':text},available,questions)

    def test_explorer_rejects_mismatched_comparisons_and_missing_combinations(self):
        import copy
        base=self.explorer_payload()
        for key,value in [('questions',{'mode':'serials','serials':'1'}),('models',['wrong']),('schemes',['original']),('averaging','schemes')]:
            body=copy.deepcopy(base);body['selections'][1][key]=value
            self.assertEqual(self.client.post('/api/explore',json=body).status_code,400)
        body=copy.deepcopy(base);body['selections'][1]['questions']={'mode':'serials','serials':'1'};body['compare']=False
        self.assertEqual(self.client.post('/api/explore',json=body).status_code,200)
        body=copy.deepcopy(base);body['selections'][0]['runs']=body['selections'][0]['runs'][:1]
        self.assertEqual(self.client.post('/api/explore',json=body).status_code,400)
        body=copy.deepcopy(base);body['pipeline']='RAG'
        self.assertEqual(self.client.post('/api/explore',json=body).status_code,400)

    def test_explorer_exports_and_local_chart_assets(self):
        from openpyxl import load_workbook
        body=self.explorer_payload(1,'schemes')
        for format_ in ('csv','xlsx'):
            response=self.client.post('/api/explore/export',json={**body,'format':format_,'table':'rows'})
            self.assertEqual(response.status_code,200,response.data)
            path=d.ARTIFACTS[response.get_json()['url'].split('/')[-1]]
            if format_=='csv':
                with path.open(encoding='utf-8-sig',newline='') as f:rows=list(csv.DictReader(f))
                self.assertEqual(len(rows),6)
                self.assertTrue(all(float(r['emissions_g'])==1 for r in rows))
            else:
                wb=load_workbook(path,read_only=True);self.assertEqual(wb['rows'].max_row,7);wb.close()
        for path in ('/ui/plotly.min.js','/ui/charts.js','/ui/app.js','/ui/styles.css'):
            with self.client.get(path) as response:self.assertEqual(response.status_code,200)
        self.assertEqual(self.client.get('/ui/dashboard.py').status_code,404)

    def test_explorer_checks_exact_text_source_changes_and_partial_coverage(self):
        body=self.explorer_payload(1)
        item=d.CATALOGUES[self.dataset][body['selections'][0]['runs'][0]]
        text=next(iter(item['run']['questions']));record=item['run']['questions'].pop(text)
        item['run']['questions'][text+' changed']=record
        self.assertEqual(self.client.post('/api/explore',json=body).status_code,400)
        item['run']['questions'].pop(text+' changed');item['run']['questions'][text]=record
        original=item['source_stamp'];item['source_stamp']=(0,0)
        self.assertEqual(self.client.post('/api/explore',json=body).status_code,400)
        item['source_stamp']=original
        # A subset may be analysed even when complete source runs contain different question sets.
        unselected=next(q for q in item['run']['questions'] if q==self.bank.iloc[19]['Question'])
        item['run']['questions'].pop(unselected)
        self.assertEqual(self.client.post('/api/explore',json=body).status_code,200)


if __name__=='__main__':unittest.main()
