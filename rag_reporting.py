"""Durable RAG-only measurement details and optional per-run Excel export."""
import csv
import json
import math
from pathlib import Path

DETAILS_FILE = 'rag_details.json'
WORKBOOK_FILE = 'modular_emissions.xlsx'
GROUP_NAMES = {
    'document_preparation': 'Document preparation',
    'embedding_indexing_storage': 'Embedding, indexing and storage',
    'retrieval': 'Question embedding, search and ranking',
    'generation': 'Context augmentation and response generation',
}
TIMING_KEYS = ('query_embedding_s', 'faiss_search_s', 'bm25_search_s', 'fusion_mmr_s', 'evidence_parent_s')


def numeric(value):
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (ValueError, TypeError):
        return None


def clean(value):
    if isinstance(value, dict):
        return {str(k): clean(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean(v) for v in value]
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if hasattr(value, 'item'):
        return clean(value.item())
    return value


def read_details(folder):
    path = Path(folder) / DETAILS_FILE
    if not path.exists():
        return None
    data = json.loads(path.read_text(encoding='utf-8'))
    if data.get('schema_version') != 1:
        raise ValueError('Unsupported RAG detail schema.')
    return data


def write_details(folder, data):
    from cancellation import protect_cleanup
    path = Path(folder) / DETAILS_FILE
    with protect_cleanup(defer_interrupt=True):
        temporary = path.with_suffix('.json.tmp')
        temporary.write_text(json.dumps(clean(data), indent=2, ensure_ascii=False, allow_nan=False), encoding='utf-8')
        temporary.replace(path)


def start_details(folder, settings, generation=None, kind='queries', export_workbook=True):
    write_details(folder, {'schema_version': 1, 'kind': kind, 'settings': dict(settings),
                           'index_generation': generation, 'export_workbook': export_workbook,
                           'status': 'running', 'queries': [], 'groups': []})


def query_entry(data, question):
    position = int(question['execution_position'])
    existing = next((q for q in data['queries'] if q['execution_position'] == position), None)
    if existing is None:
        existing = {**clean(question), 'status': 'running', 'retrieval': None}
        data['queries'].append(existing)
    return existing


def register_phase(state, tracker, phase):
    data = read_details(state['folder'])
    if data is None:
        return
    from measurement import tracker_run_id
    question = state.get('question')
    if question:
        query_entry(data, question)
    data['groups'].append({'phase': phase, 'group': GROUP_NAMES.get(phase, phase),
                           'run_id': tracker_run_id(tracker),
                           'execution_position': question.get('execution_position') if question else None,
                           'model': state.get('model'), 'status': 'running',
                           'duration_s': None, 'wall_time_s': None, 'emissions_kg': None,
                           'energy_kwh': None, 'cpu_energy_kwh': None, 'gpu_energy_kwh': None,
                           'ram_energy_kwh': None})
    write_details(state['folder'], data)


def stopped_phase(state, tracker, phase, emissions, wall_time):
    data = read_details(state['folder'])
    if data is None:
        return
    from measurement import tracker_run_id
    run_id = tracker_run_id(tracker)
    group = next((g for g in reversed(data['groups']) if g['phase'] == phase and g['run_id'] == run_id), None)
    if group is None:
        return
    if phase == 'retrieval' and state.get('retrieval_detail') and state.get('question'):
        query_entry(data, state['question'])['retrieval'] = state.pop('retrieval_detail')
    raw = []
    for filename in ('emissions.csv', 'temp_emissions.csv'):
        path = Path(state['folder']) / filename
        if path.exists():
            with path.open(encoding='utf-8-sig', newline='') as handle:
                raw.extend(csv.DictReader(handle))
    row = next((r for r in reversed(raw) if run_id and r.get('run_id') == run_id), {})
    group.update(status='measured', wall_time_s=wall_time, duration_s=numeric(row.get('duration')),
                 emissions_kg=numeric(emissions))
    for target, source in [('energy_kwh', 'energy_consumed'), ('cpu_energy_kwh', 'cpu_energy'),
                           ('gpu_energy_kwh', 'gpu_energy'), ('ram_energy_kwh', 'ram_energy')]:
        group[target] = numeric(row.get(source))
    write_details(state['folder'], data)


def record_retrieval(result):
    from cancellation import _current
    state = _current.get()
    if not state or not state.get('rag') or not state.get('question'):
        return
    state['retrieval_detail'] = {
        'accepted': result['accepted'], 'reason': result['reason'],
        'fallback_used': not result['accepted'],
        'response_mode': 'rag' if result['accepted'] else 'base_model_fallback',
        'best_dense_similarity': max((score for _, score in result['dense']), default=None),
        'selected_child_count': len(result['selected']), 'parent_count': len(result['sources']),
        'context_characters': sum(len(s['text']) for s in result['sources']),
        'selected_children': result['selected'],
        'sources': [{k: v for k, v in s.items() if k != 'text'} for s in result['sources']],
        'timings': result['timings'],
    }


def finalize_details(folder, status):
    data = read_details(folder)
    if data is None:
        return None
    path = Path(folder) / 'answers.csv'
    rows = []
    if path.exists():
        with path.open(encoding='utf-8-sig', newline='') as handle:
            rows = list(csv.DictReader(handle))
    by_position = {int(r.get('execution_position') or i): r for i, r in enumerate(rows, 1)}
    for query in data['queries']:
        row = by_position.get(query['execution_position'])
        query['status'] = row['status'] if row else ('cancelled' if status == 'cancelled' else 'error')
        query['error'] = row.get('error', '') if row else ''
    for group in data['groups']:
        if group['status'] == 'running':
            group['status'] = 'not_completed'
        row = by_position.get(group['execution_position'])
        if row and group['phase'] == 'generation':
            group['status'] = row['status']
        elif status == 'cancelled' and group is data['groups'][-1]:
            group['status'] = 'partial_cancelled'
        elif status == 'error' and group is data['groups'][-1]:
            group['status'] = 'error'
    data['status'] = status
    write_details(folder, data)
    if data.get('export_workbook'):
        try:
            return export_workbook(folder, data)
        except Exception as exc:
            print(f'WARNING: RAG measurements saved in {DETAILS_FILE}, but Excel export failed: {exc}. '
                  'Retry with python rag_reporting.py <run-folder>.')
    return None


def workbook_tables(data):
    """Flat raw-unit tables; no invented suboperation emissions or zero-filled readings."""
    queries = {q['execution_position']: q for q in data['queries']}
    headers = ['Position', 'Question ID', 'Question', 'Topic sheet', 'Question type', 'Model',
               'Measured group', 'Group status', 'Query status', 'CodeCarbon session ID',
               'Duration (s)', 'Wall time (s)', 'Emissions (kg CO2e)', 'Total energy (kWh)',
               'CPU energy (kWh)', 'GPU energy (kWh)', 'RAM energy (kWh)']
    groups = [headers]
    for group in data['groups']:
        q = queries.get(group['execution_position'], {})
        groups.append([group['execution_position'], q.get('question_id'), q.get('question'),
                       q.get('topic_sheet'), q.get('question_type'), group.get('model'), group['group'],
                       group['status'], q.get('status'), group.get('run_id'),
                       *[group.get(k) for k in ('duration_s', 'wall_time_s', 'emissions_kg', 'energy_kwh',
                                               'cpu_energy_kwh', 'gpu_energy_kwh', 'ram_energy_kwh')]])
    timing = [['Position', 'Question ID', 'Query status', 'Response mode',
               'Embedding (s)', 'FAISS (s)', 'BM25 (s)', 'Fusion / MMR (s)', 'Gate / parent expansion (s)']]
    evidence = [['Position', 'Question ID', 'Query status', 'Response mode', 'Gate reason',
                 'Best dense similarity', 'Selected children', 'Parent sources', 'Context characters', 'Source files']]
    for q in data['queries']:
        r = q.get('retrieval') or {}
        timing.append([q['execution_position'], q.get('question_id'), q['status'], r.get('response_mode'),
                       *[r.get('timings', {}).get(k) for k in TIMING_KEYS]])
        evidence.append([q['execution_position'], q.get('question_id'), q['status'], r.get('response_mode'),
                         r.get('reason'), r.get('best_dense_similarity'), r.get('selected_child_count'),
                         r.get('parent_count'), r.get('context_characters'),
                         ', '.join(dict.fromkeys(s['filename'] for s in r.get('sources', [])))])
    config = [['Setting', 'Recorded value'], ['Run kind', data['kind']], ['Run status', data['status']],
              ['Index generation', data.get('index_generation')],
              ['Missing readings', 'Blank means unavailable, not zero.'],
              ['Duration / wall time', 'Duration comes from raw CodeCarbon; wall time also includes tracker lifecycle and bookkeeping.'],
              ['Long text', 'Excel cell text is limited to 32,767 characters; rag_details.json retains the full question.'],
              ['Energy / emissions scope', 'CodeCarbon supported hardware estimates; not isolated model or wall-socket power.'],
              ['Unsupported hardware', 'A zero GPU field does not prove zero physical consumption. Check measurement_metadata.json for backend coverage.'],
              ['Step timings', 'Wall-clock timings only; no separately measured step emissions.'],
              ['Corpus costs', 'Separate indexing workbook; excluded from question totals.']]
    config.extend([[key, value if isinstance(value, (str, int, float, bool)) else json.dumps(value)]
                   for key, value in data['settings'].items()])
    sheets = {'Group measurements': groups}
    if data['kind'] == 'queries':
        sheets.update({'Step timings': timing, 'Evidence': evidence})
    sheets['Configuration'] = config
    return sheets


def export_workbook(folder, data=None):
    """Portable Python runtime export; Excel itself is not required."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    data = data or read_details(folder)
    if data is None:
        raise ValueError('This run has no recorded modular RAG measurements.')
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in workbook_tables(data).items():
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        for row in rows:
            sheet.append([ILLEGAL_CHARACTERS_RE.sub('', v)[:32767] if isinstance(v, str) else v for v in row])
        sheet.freeze_panes = 'C2' if name != 'Configuration' else 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='14394B')
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                # Question/file text must never become executable Excel formulas.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.font = Font(name='Arial', size=10, color='202F3C')
                cell.alignment = Alignment(horizontal='right' if isinstance(cell.value, (int, float)) else 'left',
                                           vertical='center', wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = '0.000000E+00'
            sheet.row_dimensions[row[0].row].height = 42 if name == 'Group measurements' else 30
        for i, header in enumerate(rows[0], 1):
            width = 22
            if header in ('Question', 'Recorded value', 'Source files'):
                width = 65
            elif header in ('Measured group', 'CodeCarbon session ID', 'Setting'):
                width = 42
            elif header in ('Position', 'Question ID'):
                width = 14
            sheet.column_dimensions[get_column_letter(i)].width = width
    output = Path(folder) / WORKBOOK_FILE
    temporary = output.with_suffix('.xlsx.tmp')
    try:
        workbook.save(temporary)
        temporary.replace(output)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    return output


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='Export recorded modular RAG measurements to Excel.')
    parser.add_argument('folder', type=Path)
    args = parser.parse_args()
    print(export_workbook(args.folder))
