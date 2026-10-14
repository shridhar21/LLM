"""Plot the six-table export without changing the workbook or original reports.

python plot_run_tables.py
python plot_run_tables.py --input obsv/run_tables.xlsx --output obsv/graphs_v2
Dependencies: python -m pip install -r requirements-report.txt
"""
import argparse
import csv
import json
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.collections import LineCollection
from matplotlib.lines import Line2D
from matplotlib.ticker import MultipleLocator, FormatStrFormatter
import numpy as np
from openpyxl import load_workbook
from question_order import load_bank

ROOT = Path(__file__).resolve().parent
RUNS = ['e1', 'e2', 'e3', 'e4', 'e5']
TOPIC_COLORS = ['#0072B2', '#D55E00', '#009E73', '#CC79A7',
                '#756800', '#56B4E9', '#B44B4B', '#686868']
RUN_COLORS = ['#0072B2', '#D55E00', '#009E73', '#CC79A7', '#756800']
HEADERS = ['Original question serial', 'Question type - topic sheet',
           'Execution position', 'CPU energy (kWh)', 'GPU energy (kWh)',
           'RAM energy (kWh)', 'Calculated CPU emissions (g CO2eq)',
           'Calculated GPU emissions (g CO2eq)', 'Calculated RAM emissions (g CO2eq)',
           'Latency (s)']


def load_data(path):
    """Use numeric stored cells, not their rounded display strings; fail on gaps."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if set(wb.sheetnames) != set(['Comprehensive', *RUNS]):
            raise ValueError('Expected Comprehensive and e1 through e5 sheets')
        records, labels = [], None
        for name in RUNS:
            ws = wb[name]
            if [c.value for c in next(ws.iter_rows(max_row=1, max_col=10))] != HEADERS:
                raise ValueError(f'{name}: unexpected column headers')
            rows = list(ws.iter_rows(min_row=2, max_row=161, max_col=10, values_only=True))
            if [r[0] for r in rows] != list(range(1, 161)):
                raise ValueError(f'{name}: expected original question serials 1..160')
            current = [r[1] for r in rows]
            if any(not isinstance(s, str) or ' - ' not in s for s in current):
                raise ValueError(f'{name}: missing question type/topic labels')
            if labels is not None and labels != current:
                raise ValueError(f'{name}: question labels do not match other runs')
            labels = current
            raw = [r[2:] for r in rows]
            if any(not isinstance(v, (float, int)) or isinstance(v, bool)
                   for row in raw for v in row):
                raise ValueError(f'{name}: missing or nonnumeric measurements; not replaced with zero')
            data = np.asarray(raw, dtype=float)
            if not np.isfinite(data).all() or (data < 0).any():
                raise ValueError(f'{name}: nonfinite or negative measurements')
            if sorted(data[:, 0]) != list(range(1, 161)):
                raise ValueError(f'{name}: execution positions are not a permutation of 1..160')
            records.append(data)
        values = np.stack(records, axis=1)  # question, run, metric
        combined = wb['Comprehensive']
        expected_headers = HEADERS[:2] + [f'{run}: {metric}' for metric in HEADERS[2:] for run in RUNS]
        if [c.value for c in next(combined.iter_rows(max_row=1, max_col=42))] != expected_headers:
            raise ValueError('Comprehensive: unexpected headers')
        for i, row in enumerate(combined.iter_rows(min_row=2, max_row=161, max_col=42, values_only=True)):
            expected = [i + 1, labels[i], *values[i].T.ravel().tolist()]
            if list(row) != expected:
                raise ValueError(f'Comprehensive disagrees with individual sheets at question {i+1}')
        types, topics = zip(*(s.rsplit(' - ', 1) for s in labels))
        return dict(position=values[:, :, 0], components=values[:, :, 1:4],
                    energy=values[:, :, 1:4].sum(axis=2), latency=values[:, :, 7],
                    types=np.array(types), topics=np.array(topics), labels=labels)
    finally:
        wb.close()


def clean(ax, grid='x'):
    ax.spines[['top', 'right']].set_visible(False)
    ax.spines[['bottom', 'left']].set_color('#C8CDD2')
    ax.set_axisbelow(True)
    ax.grid(axis=grid, color='#E8EBEE', linewidth=.6)
    ax.tick_params(length=0, pad=6)


def order_analysis(data, reports):
    """Cross-check the workbook against saved runs and original workbook identities."""
    bank = load_bank()
    if len(bank) != 160:
        raise ValueError('Original questionbank must contain 160 questions')
    questions = bank['Question'].tolist()
    if len(set(questions)) != len(questions):
        raise ValueError('Original questionbank has ambiguous duplicate question text')
    labels = [f'{r.question_type} - {r.topic_sheet}' for r in bank.itertuples()]
    if labels != data['labels']:
        raise ValueError('Workbook labels differ from original questionbank order')
    original = np.arange(1, 161)
    if not np.array_equal(data['position'][:, 0], original):
        raise ValueError('e1 does not match original questionbank order')
    displacement = np.abs(data['position'] - original[:, None]).astype(int)
    details, summary, emissions_by_run = [], [], []
    for k, name in enumerate(RUNS):
        folder = Path(reports) / name
        with (folder/'answers.csv').open(encoding='utf-8-sig', newline='') as handle:
            rows = list(csv.DictReader(handle))
        mapped = {r['question']: r for r in rows}
        if len(mapped) != len(rows) or set(mapped) != set(questions):
            raise ValueError(f'{name}: source question texts differ from the original bank')
        order = json.loads((folder/'question_order.json').read_text(encoding='utf-8'))
        entries = order['questions']
        metadata = {r['Question ID']: r for r in entries}
        if len(metadata) != len(entries) or set(metadata) != set(bank['Question ID']):
            raise ValueError(f'{name}: saved ordering identities differ from questionbank')
        source_components, source_emissions = [], []
        for i, (_, q) in enumerate(bank.iterrows()):
            record = mapped[q['Question']]
            meta = metadata[q['Question ID']]
            if (record['question_id'] != q['Question ID'] or
                int(record['execution_position']) != data['position'][i,k] or
                meta['execution_position'] != data['position'][i,k] or
                meta.get('conversation_chain', '') != q['conversation_chain'] or
                meta.get('conversation_turn', 0) != q['conversation_turn']):
                raise ValueError(f'{name}: position/chain mismatch for {q["Question ID"]}')
            component = np.array([float(record[f'{part}_energy_kwh']) for part in ['cpu','gpu','ram']])
            emission = float(record['emissions_kg'])
            if record['status'] != 'ok' or not np.isfinite(component).all() or (component < 0).any() or not np.isfinite(emission) or emission < 0:
                raise ValueError(f'{name}: invalid follow-up analysis source measurements')
            if not np.allclose(component, data['components'][i,k], rtol=1e-12, atol=1e-18):
                raise ValueError(f'{name}: source energy differs from workbook for {q["Question ID"]}')
            source_components.append(component)
            source_emissions.append(emission)
        source_components = np.array(source_components)
        source_emissions = np.array(source_emissions)
        emissions_by_run.append(source_emissions * 1000)
        chains = []
        for chain in dict.fromkeys(bank['conversation_chain']):
            if not chain:
                continue
            indexes = bank.index[bank['conversation_chain'] == chain].tolist()
            indexes.sort(key=lambda i: bank.loc[i, 'conversation_turn'])
            positions = data['position'][indexes, k].astype(int)
            gaps = np.diff(positions) - 1
            if (gaps < 0).any():
                raise ValueError(f'{name}: follow-up turn order is broken in {chain}')
            components = source_components[indexes].sum(axis=0)
            chains.append(dict(chain=chain, topic=bank.loc[indexes[0], 'topic_sheet'],
                               first_question=indexes[0]+1, positions=positions.tolist(),
                               gaps=gaps.tolist(), intervening=int(gaps.sum()),
                               turn_emissions_g=(source_emissions[indexes]*1000).tolist(),
                               components=components, energy=float(components.sum()),
                               emissions_g=float(source_emissions[indexes].sum()*1000)))
        if not chains:
            raise ValueError(f'{name}: no saved follow-up chains')
        details.append(chains)
        summary.append(dict(run=name, mean_displacement=float(displacement[:,k].mean()),
                            chains=len(chains), interleaved=sum(c['intervening']>0 for c in chains),
                            mean_intervening=float(np.mean([c['intervening'] for c in chains])),
                            components=np.sum([c['components'] for c in chains],axis=0),
                            energy=sum(c['energy'] for c in chains),
                            emissions_g=sum(c['emissions_g'] for c in chains)))
    data['emissions_g'] = np.stack(emissions_by_run, axis=1)
    return displacement, details, summary


def generate(input_path, output, reports=ROOT / 'emissions_reports'):
    data = load_data(input_path)
    _, chain_details, _ = order_analysis(data, reports)
    output = Path(output)
    if output.exists():
        raise FileExistsError(f'Choose a new output directory; already exists: {output}')
    output.mkdir(parents=True)
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 10,
                         'axes.titlesize': 12, 'axes.labelsize': 10,
                         'figure.facecolor': 'white', 'savefig.facecolor': 'white',
                         'pdf.fonttype': 42})
    pages = []
    pdf_path = output / 'run_comparison_graphs.pdf'
    with PdfPages(pdf_path) as pdf:
        def save(fig, name, title, note=''):
            fig.suptitle(title, x=.08, y=.985, ha='left', fontsize=17, weight='bold')
            if note:
                fig.text(.08, .025, note, fontsize=9, color='#555555', va='bottom')
            fig.savefig(output / f'{name}.png', dpi=180)
            pdf.savefig(fig)
            pages.append(name)
            plt.close(fig)

        topics = list(dict.fromkeys(data['topics']))
        topic_colors = dict(zip(topics, TOPIC_COLORS))
        if len(topics) > len(TOPIC_COLORS):
            raise ValueError('More topics than the supported eight topic colours')
        maximum = max(float(data['emissions_g'].max()) * 1.08, 1e-12)
        for k, run in enumerate(RUNS):
            fig, ax = plt.subplots(figsize=(22, 9))
            fig.subplots_adjust(left=.055, right=.985, top=.82, bottom=.18)
            order = np.argsort(data['position'][:, k])
            x = data['position'][order, k]
            y = data['emissions_g'][order, k]
            colors = [topic_colors[topic] for topic in data['topics'][order]]
            points = np.column_stack([x, y])
            segments, segment_colors = [], []
            for i in range(len(points)-1):
                middle = (points[i] + points[i+1]) / 2
                segments.extend([[points[i], middle], [middle, points[i+1]]])
                segment_colors.extend([colors[i], colors[i+1]])
            ax.add_collection(LineCollection(segments, colors=segment_colors, linewidths=1.5))
            ax.scatter(x, y, c=colors, s=10, linewidths=0, zorder=3)
            clean(ax, 'y')
            ax.set_xlim(.5, 160.5)
            ax.set_ylim(0, maximum)
            ax.set_xticks(x, [str(i+1) for i in order], rotation=90, fontsize=7.5)
            ax.set_xlabel('Original question number, arranged in actual execution order', labelpad=12)
            ax.set_ylabel('Emissions per question (g CO2eq)')
            fig.legend([Line2D([], [], color=topic_colors[t], linewidth=2) for t in topics],
                       topics, loc='upper center', bbox_to_anchor=(.52,.925), ncol=4, frameon=False)
            save(fig, f'01_emissions_{run}', f'{run}: question emissions in execution order',
                 'Labels are original questionbank serials (1-160). Colours identify topic sheets; colour changes halfway between neighbouring questions.')
        headers = ['e1 original\nposition'] + [f'{r}\n{label}' for r in RUNS[1:]
                   for label in ['Position', 'Gap']] + [f'{r} emissions\n(g CO2eq)' for r in RUNS]
        by_chain = [{c['chain']: c for c in chains} for chains in chain_details]
        for start in range(0, len(chain_details[0]), 2):
            fig, axes = plt.subplots(2, 1, figsize=(15, 9))
            fig.subplots_adjust(left=.035, right=.97, top=.88, bottom=.12, hspace=.45)
            for ax, base in zip(axes, chain_details[0][start:start+2]):
                ax.axis('off')
                ax.set_title(f'{base["topic"]} - follow-up set starting at Q{base["first_question"]}', loc='left', pad=12)
                chains = [run[base['chain']] for run in by_chain]
                rows = []
                for turn, original in enumerate(base['positions']):
                    rows.append([original, *[v for c in chains[1:] for v in
                                [c['positions'][turn], c['gaps'][turn-1] if turn else 'N/A']],
                                *[f'{c["turn_emissions_g"][turn]:.8g}' for c in chains]])
                table = ax.table(cellText=rows, colLabels=headers, cellLoc='center',
                                 colWidths=[.065]*9+[.083]*5, bbox=[0,0,1,1])
                table.auto_set_font_size(False); table.set_fontsize(9)
                for (row, col), cell in table.get_celld().items():
                    cell.set_facecolor('white'); cell.set_edgecolor('#D7D7D7'); cell.set_linewidth(.35)
                    if row == 0: cell.set_text_props(weight='bold')
            if len(chain_details[0][start:start+2]) < 2:
                axes[1].axis('off')
            save(fig, f'09_followup_sets_{start//2+1:02}', 'Follow-up positions, interleaving and emissions',
                 'Gap = questions outside this chain between the previous turn and this turn, including turns from other chains.\n'
                 'N/A = first turn, with no preceding turn. Emissions are for the individual question, taken from original run records.')
        # Final section: compare the same question across all five runs.
        for start in range(0, 160, 32):
            stop = min(start + 32, 160)
            values = data['emissions_g'][start:stop]
            row_positions = np.arange(stop-start)
            fig, ax = plt.subplots(figsize=(16, 10))
            fig.subplots_adjust(left=.23, right=.97, top=.86, bottom=.11)
            ax.hlines(row_positions, values.min(axis=1), values.max(axis=1),
                      color='#B4BBC1', linewidth=1.6, zorder=2)
            for k, (run, color) in enumerate(zip(RUNS, RUN_COLORS)):
                ax.scatter(values[:,k], row_positions+(k-2)*.10,
                           color=color, label=run, s=24, linewidths=0, zorder=3)
            ax.set_yticks(row_positions,
                          [f'Q{i+1}  {data["topics"][i]}' for i in range(start,stop)], fontsize=9)
            ax.set_ylim(stop-start-.5, -.5)
            ax.set_xlim(0, maximum)
            ax.xaxis.set_major_locator(MultipleLocator(.025))
            ax.xaxis.set_major_formatter(FormatStrFormatter('%.3f'))
            ax.tick_params(axis='x', labelsize=8)
            ax.set_xlabel('Emissions per question (g CO2eq)')
            clean(ax)
            fig.legend([Line2D([], [], marker='o', linestyle='', color=c) for c in RUN_COLORS],
                       RUNS, loc='upper center', bbox_to_anchor=(.6,.94), ncol=5, frameon=False)
            save(fig, f'10_question_emissions_{start//32+1:02}',
                 f'Per-question emissions across runs: Q{start+1}-Q{stop}',
                 'Original questionbank order. Each dot is one run; grey lines span minimum to maximum.\n'
                 'Small vertical offsets separate overlapping dots. All five pages use the same emissions scale.')
        # Last page: arithmetic mean for each matched question across e1-e5.
        x = np.arange(1, 161)
        average = data['emissions_g'].mean(axis=1)
        colors = [topic_colors[t] for t in data['topics']]
        points = np.column_stack([x, average])
        segments, segment_colors = [], []
        for i in range(159):
            middle = (points[i] + points[i+1]) / 2
            segments.extend([[points[i], middle], [middle, points[i+1]]])
            segment_colors.extend([colors[i], colors[i+1]])
        fig, ax = plt.subplots(figsize=(22, 9))
        fig.subplots_adjust(left=.055, right=.985, top=.82, bottom=.18)
        ax.add_collection(LineCollection(segments, colors=segment_colors, linewidths=1.5))
        ax.scatter(x, average, c=colors, s=10, linewidths=0, zorder=3)
        clean(ax, 'y')
        ax.set_xlim(.5, 160.5)
        ax.set_ylim(0, max(float(average.max())*1.08, 1e-12))
        ax.set_xticks(x, [str(i) for i in x], rotation=90, fontsize=7.5)
        ax.set_xlabel('Original questionbank number (1-160)', labelpad=12)
        ax.set_ylabel('Mean emissions per question (g CO2eq)')
        fig.legend([Line2D([], [], color=topic_colors[t], linewidth=2) for t in topics],
                   topics, loc='upper center', bbox_to_anchor=(.52,.925), ncol=4, frameon=False)
        save(fig, '11_average_question_emissions', 'Average question emissions across e1-e5',
             'Each point is the arithmetic mean of the same question\'s recorded emissions across all five runs. Original questionbank order.')
    (output / 'README.txt').write_text(
        f'Source workbook: {Path(input_path).resolve()}\n'
        'All 160 question rows and five runs are validated against the comprehensive sheet.\n'
        'Five separate landscape emissions line graphs follow actual execution order, one per run.\n'
        'Every x-axis label is an original question serial. Line/point colours identify topic sheets consistently across runs.\n'
        'Eight follow-up tables compare turn positions, chain gaps and per-question emissions.\n'
        'Final section: five landscape dot-and-range pages cover Q1-Q160 in original questionbank order.\n'
        'Last page: per-question arithmetic mean emissions across e1-e5, in original questionbank order.\n'
        'Emissions are read from original answers.csv: emissions_kg x 1000 = g CO2eq.\n'
        'Source workbook and original run files are never modified.\n'
        'No intermediate rounding; table display labels use eight significant digits.\n', encoding='utf-8')
    return pdf_path, len(pages)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=ROOT/'obsv'/'run_tables.xlsx')
    parser.add_argument('--output', type=Path, default=ROOT/'obsv'/'graphs')
    parser.add_argument('--reports', type=Path, default=ROOT/'emissions_reports', help='Folder containing saved shuffle metadata')
    args = parser.parse_args()
    try:
        path, count = generate(args.input, args.output, args.reports)
    except (ValueError, OSError) as exc:
        parser.exit(1, f'Cannot create charts: {exc}\n')
    print(f'Created {count} PNG chart/table pages and {path}')


if __name__ == '__main__':
    main()
